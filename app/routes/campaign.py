import json
import csv
import io
import re
from datetime import datetime
from pathlib import Path
from openpyxl import load_workbook

from bson import ObjectId
from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from fastapi.encoders import jsonable_encoder

from app.core.pagination import PaginationParams, page_result
from app.core.rate_limit import rate_limit
from app.core.settings import settings
from app.core.tenancy import company_id_for_user, owned_template
from app.db.connection import db
from app.models.campaign_model import CampaignCreate
from app.services.celery_app import celery_app
from app.services.storage import storage_service
from app.utils.auth import require_roles

router = APIRouter(prefix="/campaigns", tags=["Campaigns"])
MAX_CAMPAIGN_ROWS = 5000
MAX_CAMPAIGN_FILE_BYTES = 10 * 1024 * 1024


async def _create_campaign(name: str, template_id: str, company_id: str, source: dict) -> dict:
    await owned_template(template_id, company_id)
    now = datetime.utcnow()
    document = {
        "name": name, "company_id": company_id, "template_id": template_id,
        "status": "QUEUED", "total": 0, "queued": 0, "processing": 0,
        "completed": 0, "failed": 0, "progress": 0.0,
        "source": source, "created_at": now, "updated_at": now,
    }
    result = await db.campaigns.insert_one(document)
    campaign_id = str(result.inserted_id)
    celery_app.send_task("app.worker.video_worker.expand_campaign_task", args=[campaign_id], queue="campaigns")
    return {"campaign_id": campaign_id, "status": "QUEUED"}


@router.post("", status_code=202)
async def create_campaign(
    request: CampaignCreate,
    user=Depends(rate_limit("campaign-create", settings.render_rate_limit)),
):
    if user.get("role") != "company":
        raise HTTPException(403, "Company access required")
    company_id = await company_id_for_user(user)
    if not request.customer_ids:
        raise HTTPException(400, "customer_ids cannot be empty")
    return await _create_campaign(request.name, request.template_id, company_id, {"customer_ids": request.customer_ids})


@router.post("/upload", status_code=202)
async def upload_campaign(
    name: str = Form(...), template_id: str = Form(...), data_file: UploadFile = File(...),
    user=Depends(rate_limit("campaign-upload", settings.upload_rate_limit)),
):
    if user.get("role") != "company":
        raise HTTPException(403, "Company access required")
    company_id = await company_id_for_user(user)
    extension = Path(data_file.filename or "").suffix.lower()
    if extension not in {".csv", ".xlsx"}:
        raise HTTPException(400, detail={"code": "UNSUPPORTED_FORMAT", "message": "Campaign data must be CSV or XLSX"})
    content = await data_file.read(MAX_CAMPAIGN_FILE_BYTES + 1)
    if len(content) > MAX_CAMPAIGN_FILE_BYTES:
        raise HTTPException(413, detail={"code": "FILE_TOO_LARGE", "message": "Campaign file exceeds 10 MB"})
    try:
        if extension == ".csv":
            headers = next(csv.reader(io.StringIO(content.decode("utf-8-sig"))))
        else:
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            headers = [str(value or "").strip() for value in next(workbook.active.iter_rows(values_only=True))]
            workbook.close()
    except Exception as exc:
        raise HTTPException(400, detail={"code": "INVALID_CAMPAIGN_FILE", "message": "Campaign file is corrupt or empty"}) from exc
    if not {header.strip().lower() for header in headers} & {"customer_id", "id", "distributed_id", "email"}:
        raise HTTPException(400, detail={"code": "INVALID_CAMPAIGN_FILE", "message": "Campaign file requires customer_id, distributed_id, or email"})
    path = storage_service.save(content, f"{company_id}/campaigns/imports", data_file.filename or f"campaign{extension}")
    return await _create_campaign(name, template_id, company_id, {"file_path": path, "extension": extension})


@router.get("")
async def list_campaigns(
    pagination: PaginationParams = Depends(),
    status: str | None = None,
    user=Depends(require_roles("superadmin", "company")),
):
    query = {}
    if user.get("role") != "superadmin":
        query["company_id"] = await company_id_for_user(user)
    if status:
        query["status"] = status.upper()
    if pagination.search:
        query["name"] = {"$regex": re.escape(pagination.search), "$options": "i"}
    total = await db.campaigns.count_documents(query)
    sort_field, sort_direction = pagination.sort({"created_at", "updated_at", "name", "status", "progress"})
    docs = await db.campaigns.find(query, {"source": 0}).sort(sort_field, sort_direction).skip(pagination.skip).limit(pagination.limit).to_list(pagination.limit)
    for doc in docs:
        doc["id"] = str(doc.pop("_id"))
    return jsonable_encoder(page_result(docs, total, pagination))


@router.get("/{campaign_id}")
async def get_campaign(campaign_id: str, user=Depends(require_roles("superadmin", "company"))):
    if not ObjectId.is_valid(campaign_id):
        raise HTTPException(400, "Invalid campaign_id")
    query = {"_id": ObjectId(campaign_id)}
    if user.get("role") != "superadmin":
        query["company_id"] = await company_id_for_user(user)
    campaign = await db.campaigns.find_one(query, {"source": 0})
    if not campaign:
        raise HTTPException(404, "Campaign not found")
    campaign["id"] = str(campaign.pop("_id"))
    return jsonable_encoder(campaign)
