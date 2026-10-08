import logging
import os
import shutil
from copy import deepcopy
from datetime import datetime

from bson import ObjectId

from app.services.celery_app import celery_app
from app.services.render_queue import current_render_job_id, current_render_temp_dir, dispatch_render_task
from app.services.storage import ensure_media_folder, get_media_abs_path
from app.services.sync_db import sync_db
from app.services.video_renderer import render_preview
from app.services.video_renderer import render_image_preview
from app.services.output_validator import validate_generated_video
from app.core.settings import settings
from pymongo import ReturnDocument
from app.services.kokoro_tts import synthesize_and_store_media
from app.utils.placeholders import replace_placeholders
import asyncio
import csv
import uuid
import threading
from openpyxl import load_workbook

logger = logging.getLogger(__name__)


def _now() -> datetime:
    return datetime.utcnow()


def _start_job_heartbeat(job_id: str) -> threading.Event:
    stopped = threading.Event()
    def heartbeat():
        while not stopped.wait(max(1, min(30, settings.stale_job_seconds / 3))):
            try:
                sync_db.render_jobs.update_one(
                    {"job_id": job_id, "status": "PROCESSING"},
                    {"$set": {"updated_at": _now()}},
                )
            except Exception:
                pass
    threading.Thread(target=heartbeat, daemon=True).start()
    return stopped


def _task_filter(video_task_id: str | None) -> dict | None:
    if not video_task_id:
        return None
    return {"_id": ObjectId(video_task_id)}


def _set_job(job_id: str, values: dict) -> None:
    values["updated_at"] = _now()
    sync_db.render_jobs.update_one({"job_id": job_id}, {"$set": values})


def _set_task(video_task_id: str | None, values: dict) -> None:
    filt = _task_filter(video_task_id)
    if not filt:
        return
    values["updated_at"] = _now()
    sync_db.video_tasks.update_one(filt, {"$set": values})


def _campaign_transition(job: dict, previous: str, current: str) -> None:
    campaign_id = job.get("campaign_id")
    if not campaign_id or not ObjectId.is_valid(str(campaign_id)):
        return
    field = {"QUEUED": "queued", "PROCESSING": "processing", "COMPLETED": "completed", "FAILED": "failed"}
    increments = {field[previous]: -1, field[current]: 1}
    sync_db.campaigns.update_one(
        {"_id": ObjectId(str(campaign_id))},
        {"$inc": increments, "$set": {"updated_at": _now()}},
    )
    campaign = sync_db.campaigns.find_one({"_id": ObjectId(str(campaign_id))})
    if campaign:
        finished = int(campaign.get("completed", 0)) + int(campaign.get("failed", 0))
        total = max(1, int(campaign.get("total", 0)))
        status = "COMPLETED" if finished >= total else "PROCESSING"
        sync_db.campaigns.update_one(
            {"_id": campaign["_id"]},
            {"$set": {"progress": round(finished * 100 / total, 2), "status": status, "updated_at": _now()}},
        )


def _safe_cleanup_tmp(tmp_folder: str | None) -> None:
    if not tmp_folder:
        return
    tmp_path = get_media_abs_path(tmp_folder)
    media_root = os.path.abspath(os.getenv("LOCAL_MEDIA_ROOT", "./media"))
    tmp_path_abs = os.path.abspath(tmp_path)
    if os.path.commonpath([media_root, tmp_path_abs]) != media_root:
        raise ValueError("Refusing to delete temp path outside media root")
    if os.path.basename(tmp_path_abs) not in {"tmp", "temp"}:
        raise ValueError("Refusing to delete non-job temp directory")
    shutil.rmtree(tmp_path_abs, ignore_errors=True)


@celery_app.task(
    bind=True,
    name="app.worker.video_worker.render_video_task",
    autoretry_for=(),
)
def render_video_task(self, job_id: str):
    job = sync_db.render_jobs.find_one_and_update(
        {"job_id": job_id, "status": "QUEUED"},
        {"$set": {"status": "PROCESSING", "started_at": _now(), "updated_at": _now()}},
        return_document=ReturnDocument.AFTER,
    )
    if not job:
        current = sync_db.render_jobs.find_one({"job_id": job_id})
        if current and current.get("status") in {"PROCESSING", "COMPLETED", "CANCELLED"}:
            return {"job_id": job_id, "status": current["status"]}
        raise ValueError(f"Render job not found: {job_id}")

    video_task_id = job.get("video_task_id")
    folder = job.get("folder")
    tmp_folder = job.get("tmp_folder")
    output_relative = job.get("output_relative_path") or f"{folder}/output.mp4"

    ensure_media_folder(folder)
    ensure_media_folder(tmp_folder)
    output_path = get_media_abs_path(output_relative)
    tmp_path = get_media_abs_path(tmp_folder)

    token_job = current_render_job_id.set(job_id)
    token_tmp = current_render_temp_dir.set(tmp_path)

    started = job.get("started_at") or _now()
    _campaign_transition(job, "QUEUED", "PROCESSING")
    _set_job(job_id, {
        "status": "PROCESSING",
        "progress": max(int(job.get("progress") or 0), 5),
        "started_at": job.get("started_at") or started,
        "celery_task_id": self.request.id,
        "worker_hostname": self.request.hostname,
    })
    _set_task(video_task_id, {
        "status": "PROCESSING",
        "progress": max(int(job.get("progress") or 0), 5),
        "render_job_id": job_id,
    })

    heartbeat = _start_job_heartbeat(job_id)
    try:
        payload = job.get("payload") or {}
        template = deepcopy(payload.get("template") or {})
        context = deepcopy(payload.get("context") or {})
        if not template:
            raise ValueError("Render job payload is missing template data")
        render_preview(template, context, output_path)
        output_metadata = validate_generated_video(output_path)

        output_url = job.get("download_url") or f"/render-jobs/{job_id}/download"
        completed = _now()
        _set_job(job_id, {
            "status": "COMPLETED",
            "progress": 100,
            "output_url": output_url,
            "output_path": output_path,
            "completed_at": completed,
            "render_duration": (completed - started).total_seconds(),
            "output_metadata": output_metadata,
            "error": None,
            "error_details": None,
        })
        _set_task(video_task_id, {
            "status": "COMPLETED",
            "progress": 100,
            "output_url": output_url,
            "output_video_url": output_url,
            "error": None,
            "error_details": None,
        })
        logger.info("[celery_render] job_id=%s completed", job_id)
        _campaign_transition(job, "PROCESSING", "COMPLETED")
        return {"job_id": job_id, "status": "COMPLETED", "output_url": output_url}
    except Exception as exc:
        details = {
            "type": exc.__class__.__name__,
            "message": str(exc),
        }
        retry_count = int(job.get("retry_count") or 0) + 1
        if retry_count <= settings.max_render_retries:
            _set_job(job_id, {"status": "QUEUED", "retry_count": retry_count, "error": details})
            _campaign_transition(job, "PROCESSING", "QUEUED")
            raise self.retry(exc=exc, countdown=min(60, 2 ** retry_count))
        _set_job(job_id, {
            "status": "FAILED", "retry_count": retry_count,
            "error": details, "error_details": details, "failed_at": _now(),
            "render_duration": (_now() - started).total_seconds(),
        })
        _set_task(video_task_id, {
            "status": "FAILED",
            "error": str(exc),
            "error_details": details,
        })
        _campaign_transition(job, "PROCESSING", "FAILED")
        logger.exception("[celery_render] job_id=%s failed", job_id)
        raise
    finally:
        heartbeat.set()
        current_render_job_id.reset(token_job)
        current_render_temp_dir.reset(token_tmp)
        _safe_cleanup_tmp(tmp_folder)


@celery_app.task(bind=True, name="app.worker.video_worker.render_image_task", max_retries=None)
def render_image_task(self, job_id: str):
    job = sync_db.render_jobs.find_one_and_update(
        {"job_id": job_id, "status": "QUEUED"},
        {"$set": {"status": "PROCESSING", "started_at": _now(), "updated_at": _now()}},
        return_document=ReturnDocument.AFTER,
    )
    if not job:
        current = sync_db.render_jobs.find_one({"job_id": job_id})
        return {"job_id": job_id, "status": current.get("status") if current else "MISSING"}
    started = job["started_at"]
    _campaign_transition(job, "QUEUED", "PROCESSING")
    heartbeat = _start_job_heartbeat(job_id)
    token_job = current_render_job_id.set(job_id)
    token_tmp = current_render_temp_dir.set(get_media_abs_path(job["tmp_folder"]))
    try:
        payload = job.get("payload") or {}
        template = deepcopy(payload.get("template") or {})
        context = deepcopy(payload.get("context") or {})
        template_json = template.get("template_json") or template
        output_path = get_media_abs_path(job["output_relative_path"])
        ensure_media_folder(job["folder"])
        render_image_preview(
            template_json,
            context.get("customer") or {},
            context.get("company") or {},
            output_path,
        )
        completed = _now()
        output_url = job.get("download_url") or f"/render-jobs/{job_id}/download"
        _set_job(job_id, {
            "status": "COMPLETED", "progress": 100, "output_url": output_url,
            "output_path": output_path, "completed_at": completed,
            "render_duration": (completed - started).total_seconds(), "error": None,
        })
        _campaign_transition(job, "PROCESSING", "COMPLETED")
        return {"job_id": job_id, "status": "COMPLETED", "output_url": output_url}
    except Exception as exc:
        details = {"type": exc.__class__.__name__, "message": str(exc)}
        retry_count = int(job.get("retry_count") or 0) + 1
        if retry_count <= settings.max_render_retries:
            _set_job(job_id, {"status": "QUEUED", "retry_count": retry_count, "error": details})
            _campaign_transition(job, "PROCESSING", "QUEUED")
            raise self.retry(exc=exc, countdown=min(60, 2 ** retry_count))
        _set_job(job_id, {"status": "FAILED", "retry_count": retry_count, "error": details, "failed_at": _now()})
        _campaign_transition(job, "PROCESSING", "FAILED")
        raise
    finally:
        heartbeat.set()
        current_render_job_id.reset(token_job)
        current_render_temp_dir.reset(token_tmp)
        _safe_cleanup_tmp(job.get("tmp_folder"))


@celery_app.task(name="app.worker.video_worker.recover_stuck_jobs")
def recover_stuck_jobs():
    from datetime import timedelta
    cutoff = _now() - timedelta(seconds=settings.stale_job_seconds)
    requeued = 0
    failed = 0
    for job in sync_db.render_jobs.find({"status": "PROCESSING", "updated_at": {"$lt": cutoff}}):
        retry_count = int(job.get("retry_count") or 0) + 1
        if retry_count > settings.max_render_retries:
            sync_db.render_jobs.update_one(
                {"_id": job["_id"], "status": "PROCESSING"},
                {"$set": {"status": "FAILED", "retry_count": retry_count, "failed_at": _now(), "updated_at": _now(), "error": {"code": "WORKER_TIMEOUT", "message": "Worker did not update the job before its timeout"}}},
            )
            _campaign_transition(job, "PROCESSING", "FAILED")
            failed += 1
            continue
        result = sync_db.render_jobs.update_one(
            {"_id": job["_id"], "status": "PROCESSING"},
            {"$set": {"status": "QUEUED", "retry_count": retry_count, "updated_at": _now()}},
        )
        if not result.modified_count:
            continue
        payload = job.get("payload") or {}
        if job.get("job_type") == "tts":
            task_name, queue = "app.worker.video_worker.generate_tts_task", "tts"
        elif payload.get("dynamic_tts"):
            task_name, queue = "app.worker.video_worker.prepare_render_tts_task", "tts"
        elif str(job.get("output_relative_path", "")).lower().endswith((".jpg", ".jpeg", ".png", ".webp")):
            task_name, queue = "app.worker.video_worker.render_image_task", "render_image"
        else:
            task_name, queue = "app.worker.video_worker.render_video_task", "render_video"
        celery_app.send_task(task_name, args=[job["job_id"]], queue=queue)
        _campaign_transition(job, "PROCESSING", "QUEUED")
        requeued += 1
    return {"requeued": requeued, "failed": failed}


def _campaign_customer_ids(campaign: dict) -> list[str]:
    source = campaign.get("source") or {}
    if source.get("customer_ids"):
        return [str(value) for value in source["customer_ids"]][:5000]
    path = get_media_abs_path(source["file_path"])
    rows: list[dict] = []
    if source.get("extension") == ".csv":
        with open(path, "r", encoding="utf-8-sig", newline="") as handle:
            rows = list(csv.DictReader(handle))[:5000]
    else:
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook.active
        iterator = sheet.iter_rows(values_only=True)
        headers = [str(value or "").strip() for value in next(iterator)]
        rows = [dict(zip(headers, values)) for values in iterator][:5000]
        workbook.close()
    customer_ids: list[str] = []
    for row in rows:
        customer_id = row.get("customer_id") or row.get("id")
        query = {"_id": ObjectId(str(customer_id))} if customer_id and ObjectId.is_valid(str(customer_id)) else None
        if not query and row.get("distributed_id"):
            query = {"distributed_id": str(row["distributed_id"])}
        if not query and row.get("email"):
            query = {"email": str(row["email"]).lower()}
        if query:
            query["linked_company_id"] = {"$in": [campaign["company_id"], ObjectId(campaign["company_id"])]}
            customer = sync_db.customers.find_one(query, {"_id": 1})
            if customer:
                customer_ids.append(str(customer["_id"]))
    return customer_ids


@celery_app.task(name="app.worker.video_worker.expand_campaign_task")
def expand_campaign_task(campaign_id: str):
    if not ObjectId.is_valid(campaign_id):
        raise ValueError("Invalid campaign id")
    campaign = sync_db.campaigns.find_one_and_update(
        {"_id": ObjectId(campaign_id), "status": "QUEUED"},
        {"$set": {"status": "PROCESSING", "updated_at": _now()}},
        return_document=ReturnDocument.AFTER,
    )
    if not campaign:
        return {"campaign_id": campaign_id, "status": "IGNORED"}
    template = sync_db.templates.find_one({"_id": ObjectId(campaign["template_id"]), "company_id": campaign["company_id"]})
    company = sync_db.companies.find_one({"_id": ObjectId(campaign["company_id"])})
    if not template or not company:
        sync_db.campaigns.update_one({"_id": campaign["_id"]}, {"$set": {"status": "FAILED", "error": "Template or company not found"}})
        return {"campaign_id": campaign_id, "status": "FAILED"}
    try:
        customer_ids = _campaign_customer_ids(campaign)
    except Exception as exc:
        sync_db.campaigns.update_one(
            {"_id": campaign["_id"]},
            {"$set": {"status": "FAILED", "error": str(exc), "updated_at": _now()}},
        )
        return {"campaign_id": campaign_id, "status": "FAILED"}
    now = _now()
    is_image = str(template.get("type", "video")).lower() in {"img", "image"}
    created = 0
    task_dispatches: list[tuple[str, str, str]] = []
    for customer_id in customer_ids:
        customer = sync_db.customers.find_one({
            "_id": ObjectId(customer_id),
            "linked_company_id": {"$in": [campaign["company_id"], ObjectId(campaign["company_id"])]},
        })
        if not customer:
            continue
        job_id = uuid.uuid4().hex
        folder = f"render_jobs/{job_id}"
        temp_folder = f"{folder}/temp"
        output_folder = f"{folder}/output"
        ensure_media_folder(temp_folder)
        ensure_media_folder(output_folder)
        extension = "jpg" if is_image else "mp4"
        output_relative = f"{output_folder}/output.{extension}"
        customer_context = {k: (str(v) if v is not None else "") for k, v in customer.items() if k != "_id"}
        customer_context["id"] = customer_id
        company_context = {k: (str(v) if v is not None else "") for k, v in company.items() if k != "_id"}
        company_context["id"] = campaign["company_id"]
        document = {
            "job_id": job_id, "company_id": campaign["company_id"], "campaign_id": campaign_id,
            "customer_id": customer_id, "template_id": campaign["template_id"],
            "job_type": "campaign_image" if is_image else "campaign_video", "kind": "campaign_image" if is_image else "campaign_video",
            "status": "QUEUED", "progress": 0, "retry_count": 0,
            "queued_at": now, "created_at": now, "updated_at": now,
            "started_at": None, "completed_at": None, "failed_at": None,
            "render_duration": None, "output_url": None, "error": None,
            "status_url": f"/render-jobs/{job_id}", "download_url": f"/render-jobs/{job_id}/download",
            "folder": folder, "tmp_folder": temp_folder, "output_relative_path": output_relative,
            "payload": {"template": template, "context": {"customer": customer_context, "company": company_context}},
            "idempotency_key": f"campaign:{campaign_id}:customer:{customer_id}",
        }
        try:
            sync_db.render_jobs.insert_one(document)
        except Exception:
            continue
        task_name = "app.worker.video_worker.render_image_task" if is_image else "app.worker.video_worker.render_video_task"
        queue = "render_image" if is_image else "render_video"
        task_dispatches.append((task_name, queue, job_id))
        created += 1
    sync_db.campaigns.update_one(
        {"_id": campaign["_id"]},
        {"$set": {"total": created, "queued": created, "processing": 0, "completed": 0, "failed": 0, "progress": 0.0, "updated_at": _now(), "status": "PROCESSING" if created else "FAILED"}},
    )
    for task_name, queue, job_id in task_dispatches:
        try:
            celery_app.send_task(task_name, args=[job_id], queue=queue)
        except Exception as exc:
            sync_db.render_jobs.update_one(
                {"job_id": job_id, "status": "QUEUED"},
                {"$set": {"status": "FAILED", "failed_at": _now(), "updated_at": _now(), "error": {"code": "QUEUE_UNAVAILABLE", "message": str(exc)}}},
            )
            sync_db.campaigns.update_one(
                {"_id": campaign["_id"]},
                {"$inc": {"queued": -1, "failed": 1}, "$set": {"updated_at": _now()}},
            )
    return {"campaign_id": campaign_id, "jobs_created": created}


@celery_app.task(bind=True, name="app.worker.video_worker.generate_tts_task")
def generate_tts_task(self, job_id: str):
    job = sync_db.render_jobs.find_one_and_update(
        {"job_id": job_id, "status": "QUEUED"},
        {"$set": {"status": "PROCESSING", "started_at": _now(), "updated_at": _now()}},
        return_document=ReturnDocument.AFTER,
    )
    if not job:
        current = sync_db.render_jobs.find_one({"job_id": job_id})
        return {"job_id": job_id, "status": current.get("status") if current else "MISSING"}
    payload = job.get("payload") or {}
    started = job["started_at"]
    try:
        stored = asyncio.run(synthesize_and_store_media(
            company_id=job["company_id"],
            voisetext=payload["voisetext"],
            voice=payload.get("voice", "af_heart"),
            speed=float(payload.get("speed", 1.0)),
            folder_path=job["folder"],
        ))
        completed = _now()
        output_url = job.get("download_url") or f"/render-jobs/{job_id}/download"
        output_path = get_media_abs_path(stored["file_url"])
        _set_job(job_id, {
            "status": "COMPLETED", "progress": 100, "output_url": output_url,
            "output_path": output_path, "completed_at": completed,
            "render_duration": (completed - started).total_seconds(), "error": None,
        })
        return {"job_id": job_id, "status": "COMPLETED", "output_url": output_url}
    except Exception as exc:
        details = {"type": exc.__class__.__name__, "message": str(exc)}
        retry_count = int(job.get("retry_count") or 0) + 1
        if retry_count <= settings.max_render_retries:
            _set_job(job_id, {"status": "QUEUED", "retry_count": retry_count, "error": details})
            raise self.retry(exc=exc, countdown=min(60, 2 ** retry_count))
        _set_job(job_id, {"status": "FAILED", "retry_count": retry_count, "error": details, "failed_at": _now()})
        raise


@celery_app.task(bind=True, name="app.worker.video_worker.prepare_render_tts_task")
def prepare_render_tts_task(self, job_id: str):
    job = sync_db.render_jobs.find_one({"job_id": job_id, "status": "QUEUED"})
    if not job:
        return {"job_id": job_id, "status": "IGNORED"}
    payload = deepcopy(job.get("payload") or {})
    template = payload.get("template") or {}
    template_json = template.get("template_json") or {}
    track_map = (template_json.get("design") or {}).get("trackItemsMap") or {}
    context = payload.get("context") or {}
    try:
        for track_id, item in track_map.items():
            metadata = item.get("metadata") or {}
            if item.get("type") != "audio" or not metadata.get("isCustomerField"):
                continue
            if str(metadata.get("fieldPath") or "").lower() != "voiceover":
                continue
            raw_text = item.get("voisetext") or (item.get("details") or {}).get("voisetext") or ""
            resolved_text = replace_placeholders(str(raw_text), context).strip()
            if not resolved_text:
                continue
            stored = asyncio.run(synthesize_and_store_media(
                company_id=job["company_id"], voisetext=resolved_text,
                voice=str(item.get("voice") or metadata.get("voice") or "af_heart"),
                speed=max(0.5, min(2.0, float(item.get("speed") or item.get("playbackRate") or 1.0))),
                folder_path=job["folder"],
            ))
            details = item.get("details") or {}
            details["src"] = f"./media/{stored['file_url']}"
            item["details"] = details
            item["voisetext"] = resolved_text
        payload["template"] = template
        payload["dynamic_tts"] = False
        sync_db.render_jobs.update_one({"job_id": job_id, "status": "QUEUED"}, {"$set": {"payload": payload, "updated_at": _now()}})
        dispatch_render_task("app.worker.video_worker.render_video_task", job_id, "render_video")
        return {"job_id": job_id, "status": "QUEUED", "tts_prepared": True}
    except Exception as exc:
        details = {"type": exc.__class__.__name__, "message": str(exc)}
        retry_count = int(job.get("retry_count") or 0) + 1
        if retry_count <= settings.max_render_retries:
            _set_job(job_id, {"retry_count": retry_count, "error": details})
            raise self.retry(exc=exc, countdown=min(60, 2 ** retry_count))
        _set_job(job_id, {"status": "FAILED", "retry_count": retry_count, "error": details, "failed_at": _now()})
        raise
